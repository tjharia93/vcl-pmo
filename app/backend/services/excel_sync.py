from __future__ import annotations

from datetime import datetime
from pathlib import Path

from sqlalchemy.orm import Session

from .. import models

MASTER = Path("/mnt/vimit/apps/14. AI PMO/AI Operations Centre.xlsx")


def excel_status() -> str:
    return "ok" if MASTER.exists() else "missing"


def _style_sheet(ws):
    from openpyxl.styles import Font, PatternFill, Border, Side

    header_fill = PatternFill("solid", fgColor="2B3990")
    white = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    border = Border(bottom=Side(style="thin", color="CBD2E0"))
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = white
        cell.border = border
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Calibri", size=10, color="1C1C1E")
            cell.border = border


def push_project(db: Session, project_id: str) -> Path:
    from openpyxl import Workbook

    project = db.get(models.Project, project_id)
    if not project:
        raise ValueError(f"Unknown project {project_id}")
    path = Path(project.excel_path or f"/tmp/{project.name} PMO.xlsx")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    sheets = {
        "Requirements": ["ID", "Area", "Requirement", "Priority", "Status", "Owner", "UAT", "Needs Action", "Note"],
        "Issues": ["ID", "Description", "Owner", "Status", "Resolution"],
        "UAT": ["ID", "Req ID", "Description", "Tester", "Result", "Notes"],
        "Agent Log": ["ID", "Req ID", "Agent", "Status", "Notes", "TS"],
    }
    ws = wb.active
    ws.title = "Requirements"
    for title, headers in sheets.items():
        ws = wb[title] if title in wb.sheetnames else wb.create_sheet(title)
        ws.append(headers)
        if title == "Requirements":
            for r in db.query(models.Requirement).filter_by(project_id=project_id).order_by(models.Requirement.id):
                ws.append([r.id, r.area, r.requirement, r.priority, r.status, r.owner, r.uat_result, r.needs_action, r.note])
        elif title == "Issues":
            for r in db.query(models.Issue).filter_by(project_id=project_id).order_by(models.Issue.id):
                ws.append([r.id, r.description, r.owner, r.status, r.resolution])
        elif title == "UAT":
            for r in db.query(models.UATCase).filter_by(project_id=project_id).order_by(models.UATCase.id):
                ws.append([r.id, r.req_id, r.description, r.tester, r.result, r.notes])
        else:
            for r in db.query(models.AgentLog).filter_by(project_id=project_id).order_by(models.AgentLog.id.desc()).limit(200):
                ws.append([r.id, r.req_id, r.agent, r.status, r.notes, r.ts])
        _style_sheet(ws)
    wb.save(path)
    state = db.get(models.SyncState, f"excel:{project_id}") or models.SyncState(anchor=f"excel:{project_id}")
    state.last_pushed = datetime.utcnow()
    state.status = "ok"
    state.last_error = None
    db.merge(state)
    db.commit()
    return path


def push_all(db: Session) -> list[str]:
    paths = []
    for project in db.query(models.Project).all():
        try:
            paths.append(str(push_project(db, project.id)))
        except Exception as exc:
            state = db.get(models.SyncState, f"excel:{project.id}") or models.SyncState(anchor=f"excel:{project.id}")
            state.status = "error"
            state.last_error = str(exc)
            db.merge(state)
            db.commit()
    return paths


def pull_all(db: Session) -> dict:
    now = datetime.utcnow()
    for project in db.query(models.Project).all():
        state = db.get(models.SyncState, f"excel:{project.id}") or models.SyncState(anchor=f"excel:{project.id}")
        state.last_pulled = now
        state.status = "ok" if Path(project.excel_path or "").exists() else "stale"
        db.merge(state)
    db.commit()
    return {"ok": True, "mode": "mtime-check"}
